import openpyxl,pathlib,json
p=pathlib.Path('/home/magese/projects/gisagara/staking-without-gatovu-20261010/Stacking Table rev3 - 20 L-susp to A - 2026-10-09.xlsx');w=openpyxl.load_workbook(p,data_only=True)
print(w.sheetnames)
s=w.active
print('SUPPLY HEADERS',[(openpyxl.utils.get_column_letter(c),s.cell(1,c).value,s.cell(2,c).value) for c in range(40,48)])
for sn in w.sheetnames[1:]:
 t=w[sn];print(sn,t.max_row,t.max_column);print(list(t.values)[:8])
