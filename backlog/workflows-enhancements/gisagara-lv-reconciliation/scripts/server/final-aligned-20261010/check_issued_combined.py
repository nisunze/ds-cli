import openpyxl,pathlib,json
p=pathlib.Path('/home/magese/projects/gisagara/final-aligned-20261010/overall_combined_report.xlsx');w=openpyxl.load_workbook(p,data_only=True)
for s in w:
 print('SHEET',s.title,s.max_row,s.max_column)
 for i,row in enumerate(s.values,1):
  if i<=4 or any('transform' in str(v).lower() or str(v).lower()=='total' for v in row):print(i,json.dumps(row,default=str))
