import openpyxl,pathlib,re,json
r=pathlib.Path('/home/magese/projects/gisagara')
def rows(p):
 s=openpyxl.load_workbook(p,data_only=True)['LV Summary'];c={int(re.match(r'(\d+)',str(s.cell(3,i).value)).group(1)):i for i in range(1,s.max_column+1) if re.fullmatch(r'\d+\s*kVA',str(s.cell(3,i).value))};return {str(s.cell(i,1).value).lower():{k:s.cell(i,j).value or 0 for k,j in c.items()} for i in range(5,s.max_row+1) if s.cell(i,1).value and s.cell(i,1).value!='Total'}
a=rows(r/'final-aligned-20261010/overall_combined_report.xlsx');b=rows(r/'pole-restored-combined-20261010/supply-native/combined_transformer.xlsx')
for n in a:
 if a[n]!=b[n]:
  print(n,a[n],b[n]);p=json.loads((r/'pole-restored-combined-20261010'/(n+'.json')).read_text())['layers']['tr']['features'][0]['properties'];print({k:v for k,v in p.items() if 'size' in k or 'kva' in k})
