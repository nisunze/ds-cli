import openpyxl,json,pathlib,copy
r=pathlib.Path('/home/magese/projects/gisagara');o=r/'june-approved-20261010'
s=openpyxl.load_workbook(r/'staking-before-supply.xlsx').active
print('BN existing',[(i,s.cell(i,66).value,s.cell(i,67).value) for i in (2,44,380,416,408)])
req=json.loads((o/'request.json').read_text())
for item in req['transformers']:
 if item['transformer'] not in ('upgrade_mbeho','upgrade_kimana','updrade_bazenga','rugarama'):continue
 p=pathlib.Path(item['layers']);doc=json.loads(p.read_text());pr=doc['layers']['tr']['features'][0]['properties'];print(item['transformer'],pr.get('tr_size'),pr.get('plan_kva'));pr['tr_size']=pr['plan_kva'];q=o/(item['transformer']+'-effective-input.json');q.write_text(json.dumps(doc));item['layers']=str(q)
req['out_dir']=str(o/'effective-native');(o/'effective-request.json').write_text(json.dumps(req,indent=2))
