import openpyxl,json,pathlib,re
r=pathlib.Path('/home/magese/projects/gisagara');out=r/'final-aligned-20261010';w=openpyxl.load_workbook(out/'retained-supply-native/combined_transformer.xlsx',data_only=True)
assert 'Dirty Categories' not in w.sheetnames
s=w['Transformer Sizing'];rates={str(s.cell(i,6).value).upper():s.cell(i,18).value for i in range(5,s.max_row+1) if s.cell(i,6).value}
lv=w['LV Summary']; cols={int(re.match(r'(\d+)\s*kVA',lv.cell(3,c).value).group(1)):c for c in range(1,lv.max_column+1) if isinstance(lv.cell(3,c).value,str) and re.fullmatch(r'\d+\s*kVA',lv.cell(3,c).value)}
sites={lv.cell(i,1).value:{k:lv.cell(i,c).value or 0 for k,c in cols.items()} for i in range(5,lv.max_row+1) if lv.cell(i,1).value and lv.cell(i,1).value!='Total'}
targets={'UPGRADE_MBEHO':100,'UPGRADE_KIMANA':160,'UPDRADE_BAZENGA':160,'RUGARAMA':160}
print(json.dumps({'rates':{k:rates.get(k) for k in targets},'items':{k:sites.get(k) for k in targets},'total':{k:sum(x[k] for x in sites.values()) for k in cols}}))
assert len(sites)==87
for name,kva in targets.items():assert rates[name]==kva and sites[name][kva]==1,(name,rates[name],sites[name])
totals={k:sum(x[k] for x in sites.values()) for k in cols};assert sum(totals.values())==65
assert totals=={50:20,100:27,160:15,200:1,250:2}
(out/'combined-verification.json').write_text(json.dumps({'sites':87,'supply':65,'sizes':totals,'approved_rates':targets,'dirty_categories':0,'native_output_unpatched':True},indent=2))

assert sum(sites['KIRWA'].values())==0



