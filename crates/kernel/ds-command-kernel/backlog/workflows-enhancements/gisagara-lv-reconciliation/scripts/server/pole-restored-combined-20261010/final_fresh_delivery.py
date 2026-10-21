import pathlib,json,zipfile,xml.etree.ElementTree as E,hashlib,re,struct
import openpyxl
r=pathlib.Path('/home/magese/projects/gisagara/pole-restored-combined-20261010');src=r/'planned-supply-native/combined_transformer.xlsx'
w=openpyxl.load_workbook(src,data_only=True);assert 'Dirty Categories' not in w.sheetnames
s=w['LV Summary'];cols={int(re.match(r'(\d+)',str(s.cell(3,c).value)).group(1)):c for c in range(1,s.max_column+1) if re.fullmatch(r'\d+\s*kVA',str(s.cell(3,c).value))}
sites={str(s.cell(i,1).value).upper():{k:s.cell(i,c).value or 0 for k,c in cols.items()} for i in range(5,s.max_row+1) if s.cell(i,1).value and s.cell(i,1).value!='Total'}
assert len(sites)==87
rates={'UPGRADE_MBEHO':100,'UPGRADE_KIMANA':160,'UPDRADE_BAZENGA':160,'RUGARAMA':160}
for n,k in rates.items():assert sites[n][k]==1
assert sum(sites['KIRWA'].values())==0
totals={k:sum(x[k] for x in sites.values()) for k in cols};assert sum(totals.values())==65,totals
# Remove only the requested sheet; preserve every retained worksheet byte.
with zipfile.ZipFile(src) as z:files={n:z.read(n) for n in z.namelist()}
ns='http://schemas.openxmlformats.org/spreadsheetml/2006/main';rel='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
wb=E.fromstring(files['xl/workbook.xml']);sheets=wb.find('{'+ns+'}sheets');sheet=next(x for x in sheets if x.get('name')=='Transformer Sizing');rid=sheet.get('{'+rel+'}id');removed_index=list(sheets).index(sheet);sheets.remove(sheet)
dn=wb.find('{'+ns+'}definedNames')
if dn is not None:
 for x in list(dn):
  if x.get('localSheetId')==str(removed_index):dn.remove(x)
rels=E.fromstring(files['xl/_rels/workbook.xml.rels']);rr=next(x for x in rels if x.get('Id')==rid);target=rr.get('Target');target=target.lstrip('/') if target.startswith('/') else 'xl/'+target;rels.remove(rr)
ct=E.fromstring(files['[Content_Types].xml'])
for x in list(ct):
 if x.get('PartName')=='/'+target:ct.remove(x)
files.pop(target)
for key,tree in [('xl/workbook.xml',wb),('xl/_rels/workbook.xml.rels',rels),('[Content_Types].xml',ct)]:files[key]=E.tostring(tree,encoding='utf-8',xml_declaration=True)
dst=r/'overall_combined_report.xlsx'
with zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in files.items():z.writestr(n,b)
after=openpyxl.load_workbook(dst,data_only=True);assert after.sheetnames==['InfoTable','LV Summary']
with zipfile.ZipFile(src) as a,zipfile.ZipFile(dst) as b:
 for n in b.namelist():
  if n.startswith('xl/worksheets/'):assert a.read(n)==b.read(n)
sh=r/'Shapefiles';sh.mkdir(exist_ok=True)
with zipfile.ZipFile(r/'planned-native/combined_transformer.shp.zip') as z:
 for m in z.infolist():
  if not m.is_dir():(sh/pathlib.Path(m.filename).name).write_bytes(z.read(m))
for p in sh.glob('*.shp'):
 for ext in ['.shx','.dbf','.prj']:assert p.with_suffix(ext).exists()
def dbf(p):
 b=p.read_bytes();num=struct.unpack_from('<I',b,4)[0];head,length=struct.unpack_from('<HH',b,8);fields=[];pos=32
 while b[pos]!=13:
  fields.append((b[pos:pos+11].split(b'\0')[0].decode(),b[pos+16]));pos+=32
 rows=[]
 for i in range(num):
  raw=b[head+i*length:head+(i+1)*length];off=1;row={}
  for name,size in fields:row[name]=raw[off:off+size].decode('utf-8',errors='replace').strip();off+=size
  rows.append(row)
 return rows
trp=next(sh.glob('*tr.dbf'));tr=dbf(trp);assert len(tr)==87,(trp,len(tr));print('GIS transformer fields',list(tr[0]));print([(x['names'],x['transfo'],x['plan_kva'],x['tr_size']) for x in tr if any(k in str(x).lower() for k in ['mbeho','kimana','bazeng','rugarama','kirwa'])])
for n,k in rates.items():
 row=next(x for x in tr if x.get('transfo','').upper()==n);assert float(row['plan_kva'])==k and float(row['tr_size'])==k
kirwa=next(x for x in tr if x.get('transfo','').upper()=='KIRWA');assert float(kirwa['plan_kva'])==100 and float(kirwa['tr_size'])==100
proof={'scope_sites':87,'approved_plan_kva_preserved':rates,'supply_total':65,'supply_by_kva':totals,'kirwa_supply':0,'removed_sheet':'Transformer Sizing','remaining_sheets':after.sheetnames,'retained_worksheet_bytes_unchanged':True,'shapefile_sets':len(list(sh.glob('*.shp'))),'native_GIS_saved_ratings_preserved':True,'bridges':['6b369e94-89f2-43a0-a8c5-74b8fab0dcf7','feae1d8e-271a-4a27-96cc-459b53c50ba5']}
(r/'delivery-proof.json').write_text(json.dumps(proof,indent=2))
paths=[dst]+list(sh.iterdir());(r/'delivery-hashes.json').write_text(json.dumps({str(p.relative_to(r)):hashlib.sha256(p.read_bytes()).hexdigest() for p in paths},indent=2))
with zipfile.ZipFile(r/'delivery.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in paths:z.write(p,str(p.relative_to(r)))
print(json.dumps(proof))
